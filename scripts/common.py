from __future__ import annotations

import hashlib
import os
import re
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable, Iterator

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


PMID_BLOCK = re.compile(
    r"\(\s*PMIDs?\s*[:\uFF1A#]?\s*(?P<ids>\d+(?:\s*[,;\uFF0C\uFF1B]\s*\d+)*)\s*\)",
    re.IGNORECASE,
)
PMID_SPLIT = re.compile(r"\s*[,;\uFF0C\uFF1B]\s*")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def collision_safe_path(path: Path) -> Path:
    if not path.exists():
        return path
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    candidate = path.with_name(f"{path.stem}-{timestamp}{path.suffix}")
    counter = 2
    while candidate.exists():
        candidate = path.with_name(f"{path.stem}-{timestamp}-{counter}{path.suffix}")
        counter += 1
    return candidate


def reveal_file(path: Path) -> None:
    if os.name != "nt":
        return
    subprocess.Popen(["explorer.exe", f"/select,{path.resolve()}"], close_fds=True)


def open_file(path: Path) -> None:
    if os.name == "nt":
        os.startfile(path.resolve())


def parse_pmid_block(match: re.Match[str]) -> list[str]:
    return [value for value in PMID_SPLIT.split(match.group("ids")) if value]


def iter_paragraphs(parent: DocumentObject | _Cell) -> Iterator[Paragraph]:
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        yield from iter_table_paragraphs(table)


def iter_table_paragraphs(table: Table) -> Iterator[Paragraph]:
    seen: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            identity = id(cell._tc)
            if identity in seen:
                continue
            seen.add(identity)
            yield from iter_paragraphs(cell)


def scan_groups(document: DocumentObject) -> list[list[str]]:
    groups: list[list[str]] = []
    for paragraph in iter_paragraphs(document):
        for match in PMID_BLOCK.finditer(paragraph.text):
            groups.append(parse_pmid_block(match))
    return groups


def unique_pmids(groups: Iterable[Iterable[str]]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for group in groups:
        for pmid in group:
            if pmid not in seen:
                seen.add(pmid)
                result.append(pmid)
    return result


def replace_labeled_pmids(
    document: DocumentObject,
    citation_for_pmids: Callable[[list[str]], str],
) -> int:
    replacements = 0
    for paragraph in iter_paragraphs(document):
        text = "".join(run.text for run in paragraph.runs)
        matches = list(PMID_BLOCK.finditer(text))
        for match in reversed(matches):
            replacement = citation_for_pmids(parse_pmid_block(match))
            _replace_run_span(paragraph, match.start(), match.end(), replacement)
            replacements += 1
    return replacements


def _replace_run_span(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    runs = paragraph.runs
    if not runs:
        raise RuntimeError("A PMID marker was found in a paragraph without editable runs")

    positions: list[tuple[int, int]] = []
    cursor = 0
    for run in runs:
        positions.append((cursor, cursor + len(run.text)))
        cursor += len(run.text)

    start_index = end_index = None
    start_offset = end_offset = 0
    for index, (left, right) in enumerate(positions):
        if start_index is None and left <= start < right:
            start_index = index
            start_offset = start - left
        if left < end <= right:
            end_index = index
            end_offset = end - left
            break

    if start_index is None or end_index is None:
        raise RuntimeError("Could not map a PMID marker back to Word runs")

    if start_index == end_index:
        run = runs[start_index]
        run.text = run.text[:start_offset] + replacement + run.text[end_offset:]
        return

    prefix = runs[start_index].text[:start_offset]
    suffix = runs[end_index].text[end_offset:]
    runs[start_index].text = prefix + replacement + suffix
    for index in range(start_index + 1, end_index + 1):
        runs[index].text = ""


def load_docx(path: Path) -> DocumentObject:
    if path.suffix.lower() != ".docx" or not path.is_file():
        raise ValueError(f"Input must be an existing .docx file: {path}")
    return Document(path)
